import json
import re
import os
from datetime import datetime, timedelta
from RPA.Browser.Selenium import Selenium
from RPA.Tables import Tables
from RPA.Robocorp.WorkItems import WorkItems
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

from openpyxl import Workbook
import shutil
import time

class NewsScraper:
    def __init__(self):
        self.browser = Selenium()
        self.tables = Tables()
        self.work_items = WorkItems()
        self.config = self.load_config()

        # Tentar carregar os parâmetros do work item, caso contrário usar config padrão
        try:
            self.work_items.get_input_work_item()
            self.search_phrase = self.work_items.get_work_item_variable('search_phrase')
            self.news_category = self.work_items.get_work_item_variable('news_category')
            self.months = int(self.work_items.get_work_item_variable('months'))
        except Exception as e:
            print(f"Error loading work item: {e}, using default configuration.")
            self.search_phrase = self.config['search_phrase']
            self.news_category = self.config['news_category']
            self.months = self.config['months']

        # Crie o diretório para salvar as imagens
        self.images_directory = self.create_images_directory()

    def load_config(self):
        config_path = os.path.join(os.path.dirname(__file__), 'config.json')
        if not os.path.exists(config_path):
            raise FileNotFoundError(f"Configuration file not found: {config_path}")
        with open(config_path) as config_file:
            return json.load(config_file)
        
    def create_images_directory(self):
        today = datetime.today().strftime('%Y-%m-%d')
        directory_name = f"{today}_{self.search_phrase.replace(' ', '_')}"
        directory_path = os.path.join(os.path.dirname(__file__), 'images', directory_name)
        os.makedirs(directory_path, exist_ok=True)
        return directory_path

    def open_site(self, url):
        self.browser.open_available_browser(url)

    def find_element_dynamically(self, selectors):
        for selector in selectors:
            try:
                print(f"Trying selector: {selector}")  # Log para cada seletor tentado
                element = self.browser.find_element(selector)
                print(f"Element found using selector: {selector}")
                return element
            except Exception:
                continue
        raise Exception(f"None of the selectors worked: {selectors}")

    def find_element_dynamically_in_element(self, parent, selectors):
        for selector in selectors:
            try:
                element = parent.find_element(selector)
                return element
            except Exception:
                continue
        raise Exception(f"None of the selectors worked within the element: {selectors}")
    

    def close_cookies_banner(self):
        # Tentar fechar o banner de cookies, se presente
        try:
            self.browser.execute_javascript("""
                const cookiesButtons = document.querySelectorAll('button, div[role="button"], span[role="button"]');
                for (const button of cookiesButtons) {
                    if (button.textContent.toLowerCase().includes('accept') || button.textContent.toLowerCase().includes('agree')) {
                        button.click();
                        break;
                    }
                }
            """)
            print("Cookies banner closed.")
        except Exception as e:
            print(f"Failed to close cookies banner: {e}")

    def search_news(self):

        time.sleep(10)

        # Tentar fechar o banner de cookies, se presente
        self.close_cookies_banner()


        # Primeiro, encontrar e clicar no botão que contém o ícone de busca (lupa)
        time.sleep(10)

        # Primeiro, encontrar e clicar no botão que contém o ícone de busca (lupa)
        button_found = self.browser.execute_javascript("""
            const logs = [];
            logs.push('Searching for SVGs...');
            const svgs = document.querySelectorAll('svg');
            logs.push('SVGs found: ' + svgs.length);
            
            // Procurar por ícones de pesquisa em SVGs e seus pais
            for (const svg of svgs) {
                const attributes = svg.attributes;
                let isSearchIcon = false;
                for (const attr of attributes) {
                    logs.push('SVG attribute: ' + attr.name + ', ' + attr.value);
                    if (attr.value.includes('search')) {
                        isSearchIcon = true;
                        break;
                    }
                }
                if (isSearchIcon) {
                    let button = svg.closest('button, a, div[role="button"], span[role="button"]');
                    if (button) {
                        logs.push('Found search icon, clicking button...');
                        const event = new MouseEvent('click', {
                            view: window,
                            bubbles: true,
                            cancelable: true
                        });
                        button.dispatchEvent(event);
                        return {logs: logs, clicked: true};
                    }
                }
            }
            
            // Se nenhum SVG foi identificado como ícone de busca, procurar o botão usando o atributo data-testid
            logs.push('Searching for button with data-testid...');
            let button = document.querySelector('button[data-testid="search-button"]');
            if (button) {
                logs.push('Found search button, attempting to click...');
                const event = new MouseEvent('click', {
                    view: window,
                    bubbles: true,
                    cancelable: true
                });
                button.dispatchEvent(event);

                // Verificando se o botão está visível e habilitado
                if (button.offsetParent !== null && !button.disabled) {
                    logs.push('Button is visible and enabled, clicked successfully.');
                    return {logs: logs, clicked: true};
                } else {
                    logs.push('Button is not visible or enabled.');
                }
            } else {
                logs.push('Search button not found.');
            }

            return {logs: logs, clicked: false};
        """)

        for log in button_found['logs']:
            print(log)

        if not button_found['clicked']:
            raise Exception("Search button not found or not interactable")
        
        time.sleep(10)

        # Usar JavaScript para encontrar dinamicamente a barra de pesquisa
        search_field = self.browser.execute_javascript("""
            function findSearchField() {
                const inputs = document.querySelectorAll('input');
                for (const input of inputs) {
                    if (input.placeholder && input.placeholder.toLowerCase().includes('search')) {
                        return input;
                    }
                    if (input.title && input.title.toLowerCase().includes('search')) {
                        return input;
                    }
                }
                return null;
            }
            return findSearchField();
        """)

        if search_field is None:
            raise Exception("Search field not found")

        time.sleep(10)
        self.browser.input_text(search_field, self.search_phrase)
        time.sleep(10)
        self.browser.press_keys(search_field, 'ENTER')
        time.sleep(10)

    def filter_by_category(self):
        if self.news_category:
            category_selectors = [
                f'css:[data-category="{self.news_category}"]', 
                f'css:.category-{self.news_category}', 
                f'css:a[href*="{self.news_category}"]'
            ]
            category_element = self.find_element_dynamically(category_selectors)
            
            # Adicionar espera explícita para garantir que o elemento esteja visível e clicável
            WebDriverWait(self.browser.driver, 10).until(
                EC.element_to_be_clickable((By.CSS_SELECTOR, category_element))
            )
            
            self.browser.click_element(category_element)

    def sort_results_by_date(self):
        time.sleep(10)

        # Selecionar a opção "date" no dropdown de ordenação
        try:
            self.browser.select_from_list_by_value('id:search-sort-option', 'date')
            print("Sorted results by date.")
        except Exception as e:
            print(f"Failed to sort results by date: {e}")

        time.sleep(10)

    def close_registration_popup(self):
        # Tentar fechar o popup de registro, se presente
        try:
            self.browser.execute_javascript("""
                const popupCloseButton = document.querySelector('button.close-button');
                if (popupCloseButton) {
                    popupCloseButton.click();
                }
            """)
            print("Registration popup closed.")
        except Exception as e:
            print(f"Failed to close registration popup: {e}")

    def save_image(self, image_url, file_path):
        file_url = f"file:///{file_path.replace(os.sep, '/')}"
        self.browser.execute_javascript(f"""
        var link = document.createElement('a');            
        link.href = '{image_url}';            
        link.download = '{file_url}';            
        document.body.appendChild(link);            
        link.click();            
        document.body.removeChild(link);
        """)

        # Mover a imagem da pasta de downloads para a pasta especificada
        downloads_folder = os.path.join(os.path.expanduser('~'), 'Downloads')
        downloaded_file = os.path.join(downloads_folder, os.path.basename(image_url))
        time.sleep(5)  # Aguarde o download ser concluído

        if os.path.exists(downloaded_file):
            shutil.move(downloaded_file, file_path)
            print(f"Image moved to {file_path}")
        else:
            print(f"Downloaded file not found: {downloaded_file}")
        

    def extract_news(self):
        news_data = []

        time.sleep(20)
        self.close_registration_popup()

        time.sleep(10)

        # Seletores específicos para o card de notícias da BBC
        news_items = self.browser.find_elements('css:div[data-testid="liverpool-article"]')

        for item in news_items:
                try:
                    print(item)  # Debugging: Print the WebElement

                    title_element = item.find_element('css selector', 'h2[data-testid="card-headline"]')
                    title = title_element.text if title_element else "N/A"

                    date_element = item.find_element('css selector', 'span[data-testid="card-metadata-lastupdated"]')
                    date = date_element.text if date_element else "N/A"

                    description_element = item.find_element('css selector', 'p[data-testid="card-description"]')
                    description = description_element.text if description_element else "N/A"

                    image_element = item.find_element('css selector', 'div[data-testid="card-media"] img')
                    image_url = image_element.get_attribute('src') if image_element else "N/A"
                    image_filename = os.path.basename(image_url) if image_url else "N/A"

                    news_data.append({
                        'title': title,
                        'date': date,
                        'description': description,
                        'image_filename': image_filename,
                        'count_search_phrase': title.count(self.search_phrase) + description.count(self.search_phrase),
                        'contains_money': bool(re.search(r'\$\d+(?:,\d{3})*(?:\.\d{2})?|dollars|USD', title + description))
                    })

                    # Download the image if URL is valid
                    if image_url and image_url != "N/A":
                        time.sleep(5)
                        self.save_image(image_url, os.path.join(self.images_directory, image_filename))
                        time.sleep(5)
                except Exception as e:
                    print(f"Error extracting news item: {e}")
                    continue

        for x in news_data:
            print(x)

        return news_data

    #def save_to_excel(self, data):
    #    self.tables.create_table(data)
    #    self.tables.save_table('news_data.xlsx')
    #    print("Data saved to news_data.xlsx")

    #def save_to_excel(self, news_data):
    #    self.tables.create_table('news_data', news_data)
    #    self.tables.save_workbook('news_data.xlsx')

    def find_elements_dynamically(self, selectors):
        for selector in selectors:
            try:
                print(f"Trying selector: {selector}")
                elements = self.browser.find_elements(selector.split(':')[1])
                if elements:
                    print(f"Found elements with selector: {selector}")
                    return elements
            except Exception as e:
                print(f"Error finding elements with selector {selector}: {e}")
                continue
        raise Exception("No elements found using provided selectors")

    def find_element_by_keywords(self, parent_element, keywords):
        for keyword in keywords:
            for tag in ['div', 'span', 'p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'time', 'img']:
                selector = f'{tag}[class*="{keyword}"], {tag}[id*="{keyword}"], {tag}[name*="{keyword}"], {tag}[data-*="{keyword}"]'
                try:
                    print(f"Trying keyword: {keyword} with tag: {tag}")
                    element = parent_element.find_element_by_css_selector(selector)
                    if element:
                        print(f"Found element with keyword: {keyword} and tag: {tag}")
                        return element
                except Exception as e:
                    print(f"Error finding element with keyword {keyword}: {e}")
                    continue
        raise Exception(f"Element not found using keywords: {keywords}")


    def save_to_excel(self, news_data):
        # Garantir que o formato dos dados esteja correto
        if isinstance(news_data, list) and all(isinstance(item, dict) for item in news_data):
            # Verificar se todos os dicionários possuem as mesmas chaves
            all_keys = set().union(*(d.keys() for d in news_data))
            for item in news_data:
                for key in all_keys:
                    if key not in item:
                        item[key] = "N/A"
            
            # Criar um novo workbook e uma nova planilha
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "News Data"
            
            # Adicionar os cabeçalhos na primeira linha
            headers = list(all_keys)
            sheet.append(headers)
            
            # Adicionar os dados das notícias
            for item in news_data:
                row = [item.get(key, "N/A") for key in headers]
                sheet.append(row)
            
            # Definir o caminho para salvar o arquivo Excel
            excel_path = os.path.join(os.path.dirname(__file__), 'news_data.xlsx')
            workbook.save(excel_path)
            print(f"Data saved to {excel_path}")
        else:
            raise TypeError("Not a valid input format")

    def close(self):
        self.browser.close_all_browsers()